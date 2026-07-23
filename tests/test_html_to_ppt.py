"""Tests for the HTML -> PowerPoint converter."""

import os

from pptx import Presentation

from src import html_to_ppt


def _slide_text(path: str) -> str:
    prs = Presentation(path)
    chunks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
    return "\n".join(chunks)


def test_headings_become_slides(tmp_path):
    html = "<body><h1>First</h1><p>alpha</p><h2>Second</h2><p>beta</p></body>"
    out = str(tmp_path / "deck.pptx")
    html_to_ppt.html_to_pptx(html, out)

    prs = Presentation(out)
    assert len(prs.slides) == 2
    text = _slide_text(out)
    assert "First" in text and "alpha" in text
    assert "Second" in text and "beta" in text


def test_list_items_are_bulleted(tmp_path):
    html = "<body><h1>T</h1><ul><li>one</li><li>two</li></ul></body>"
    out = str(tmp_path / "deck.pptx")
    html_to_ppt.html_to_pptx(html, out)
    text = _slide_text(out)
    assert "• one" in text and "• two" in text


def test_flat_html_falls_back_to_text(tmp_path):
    html = "just some words with no structure at all"
    out = str(tmp_path / "deck.pptx")
    html_to_ppt.html_to_pptx(html, out)
    assert len(Presentation(out).slides) >= 1


def test_returns_a_real_file_when_no_path_given():
    path = html_to_ppt.html_to_pptx("<body><h1>Hi</h1></body>")
    try:
        assert path.endswith(".pptx")
        assert os.path.getsize(path) > 0
    finally:
        os.unlink(path)
